from __future__ import annotations

import os
from dataclasses import replace

from docx import Document

import secure_rag.indexer as indexer_module
import secure_rag.retrieval as retrieval_module
from secure_rag.config import AppConfig, load_config
from secure_rag.embeddings import HashingEmbedder
from secure_rag.events import PipelineEvent
from secure_rag.extracted_cache import (
    ExtractedTextCache,
    ProcessMemoryExtractedTextCache,
)
from secure_rag.extractors import ExtractedDocument
from secure_rag.extractors import extract_document as canonical_extract
from secure_rag.extractors import file_sha256 as canonical_file_sha256
from secure_rag.indexer import Indexer
from secure_rag.manifest import ManifestStore
from secure_rag.retrieval import Retriever
from secure_rag.vector_store import QdrantStore


def _test_config(tmp_path, *, persist_cache: bool = True) -> AppConfig:
    source = tmp_path / "source"
    source.mkdir()
    base = load_config()
    config = replace(
        base,
        paths=replace(
            base.paths,
            source_root=source.resolve(),
            runtime_root=(tmp_path / "runtime").resolve(),
        ),
        ingestion=replace(
            base.ingestion,
            persist_extracted_text_cache=persist_cache,
        ),
    )
    config.ensure_runtime()
    return config


def test_process_memory_cache_is_policy_keyed_and_strictly_bounded(tmp_path) -> None:
    config = _test_config(tmp_path, persist_cache=False)
    cache = ProcessMemoryExtractedTextCache(
        max_entries=2,
        max_total_chars=8,
        max_entry_chars=6,
    )
    revision = "f" * 64
    first = ExtractedDocument("aaaa")
    second = ExtractedDocument("bbbb")
    third = ExtractedDocument("cccc")

    assert cache.store("1" * 32, revision, ".txt", config.ingestion, first)
    assert cache.store("2" * 32, revision, ".txt", config.ingestion, second)
    assert cache.entry_count == 2
    assert cache.total_chars == 8

    # Touch the first item: the second one becomes the least recently used entry.
    assert cache.load("1" * 32, revision, ".txt", config.ingestion) == first
    assert cache.store("3" * 32, revision, ".txt", config.ingestion, third)
    assert cache.load("2" * 32, revision, ".txt", config.ingestion) is None
    assert cache.load("1" * 32, revision, ".txt", config.ingestion) == first
    assert cache.load("3" * 32, revision, ".txt", config.ingestion) == third
    assert cache.entry_count == 2
    assert cache.total_chars == 8

    # Oversized entries are skipped and can never force the configured bounds higher.
    assert not cache.store(
        "4" * 32,
        revision,
        ".txt",
        config.ingestion,
        ExtractedDocument("1234567"),
    )
    assert cache.entry_count == 2
    assert cache.total_chars == 8

    changed_policy = replace(
        config.ingestion,
        max_file_mb=config.ingestion.max_file_mb + 1,
    )
    assert cache.load("1" * 32, revision, ".txt", changed_policy) is None


def test_streamlit_style_ram_cache_avoids_extraction_without_disk_cache(
    tmp_path, monkeypatch
) -> None:
    config = _test_config(tmp_path, persist_cache=False)
    source_file = config.paths.source_root / "report.docx"
    source_document = Document()
    source_document.add_paragraph(
        "Испытание турбины показало вибрацию на расчетном режиме."
    )
    source_document.save(source_file)

    embedder = HashingEmbedder(64)
    memory_cache = ProcessMemoryExtractedTextCache(
        max_entries=4,
        max_total_chars=20_000,
        max_entry_chars=10_000,
    )
    persistent_root = config.paths.runtime_root / "extracted-text-cache"
    with (
        ManifestStore(config.manifest_path) as manifest,
        QdrantStore(config.qdrant_path, "memory_cache_test", embedder.dimension) as store,
    ):
        assert Indexer(config, embedder, manifest, store).run().indexed == 1
        assert not persistent_root.exists()
        assert Retriever(config, embedder, manifest, store).extracted_cache is None

        first_events: list[PipelineEvent] = []
        assert len(
            Retriever(
                config,
                embedder,
                manifest,
                store,
                extracted_cache=memory_cache,
            ).search("вибрация", on_event=first_events.append)
        ) == 1
        assert memory_cache.entry_count == 1
        assert not persistent_root.exists()

        def unexpected_extraction(*_args, **_kwargs):
            raise AssertionError("RAM cache hit must avoid repeated extraction")

        source_hashes = 0

        def counted_file_sha256(*args, **kwargs):
            nonlocal source_hashes
            source_hashes += 1
            return canonical_file_sha256(*args, **kwargs)

        monkeypatch.setattr(retrieval_module, "extract_document", unexpected_extraction)
        monkeypatch.setattr(retrieval_module, "file_sha256", counted_file_sha256)
        second_events: list[PipelineEvent] = []
        assert len(
            Retriever(
                config,
                embedder,
                manifest,
                store,
                extracted_cache=memory_cache,
            ).search("вибрация", on_event=second_events.append)
        ) == 1
        assert source_hashes == 1
        completed = next(
            event
            for event in second_events
            if event.stage == "retrieval.source_read" and event.status == "completed"
        )
        assert completed.details["source_hashes"] == 1
        assert completed.details["extracted_cache_hits"] == 1
        assert completed.details["extractions"] == 0
        assert not persistent_root.exists()


def test_cache_policy_and_corruption_invalidate_entry(tmp_path) -> None:
    config = _test_config(tmp_path)
    cache = ExtractedTextCache(config.paths.runtime_root / "extracted-text-cache")
    document_id = "a" * 32
    revision = "b" * 64
    extracted = ExtractedDocument("Нормализованный текст", page_count=2, truncated=False)

    path = cache.store(document_id, revision, ".pdf", config.ingestion, extracted)
    assert cache.load(document_id, revision, ".pdf", config.ingestion) == extracted
    assert list(path.parent.glob("*.tmp")) == []

    changed_policy = replace(
        config.ingestion,
        max_extracted_chars=config.ingestion.max_extracted_chars + 1,
    )
    assert cache.load(document_id, revision, ".pdf", changed_policy) is None
    assert not path.exists()

    cache.store(document_id, revision, ".pdf", config.ingestion, extracted)
    path.write_bytes(b"corrupted-cache-without-source-text")
    assert cache.load(document_id, revision, ".pdf", config.ingestion) is None
    assert not path.exists()


def test_retrieval_uses_cache_and_rejects_source_tamper(tmp_path, monkeypatch) -> None:
    config = _test_config(tmp_path)
    source_file = config.paths.source_root / "report.docx"
    source_document = Document()
    source_document.add_paragraph(
        "Испытание турбины показало вибрацию на расчетном режиме."
    )
    source_document.save(source_file)

    embedder = HashingEmbedder(64)
    cache = ExtractedTextCache(config.paths.runtime_root / "extracted-text-cache")
    with (
        ManifestStore(config.manifest_path) as manifest,
        QdrantStore(config.qdrant_path, "cache_test", embedder.dimension) as store,
    ):
        report = Indexer(config, embedder, manifest, store).run()
        assert report.indexed == 1
        record = manifest.all_documents()[0]
        cache_path = cache.path_for(record.document_id, record.revision)
        assert cache_path.is_file()

        def unexpected_extraction(*_args, **_kwargs):
            raise AssertionError("cache hit must not extract the DOCX again")

        monkeypatch.setattr(retrieval_module, "extract_document", unexpected_extraction)
        events: list[PipelineEvent] = []
        hits = Retriever(config, embedder, manifest, store).search(
            "вибрация турбины", on_event=events.append
        )
        assert len(hits) == 1
        completed = next(
            event
            for event in events
            if event.stage == "retrieval.source_read" and event.status == "completed"
        )
        assert completed.details["extracted_cache_hits"] == 1
        assert completed.details["extractions"] == 0

        # A cache corruption is a miss: regenerate from the still-verified source.
        cache_path.write_bytes(b"invalid-gzip")
        calls = 0

        def counted_extraction(*args, **kwargs):
            nonlocal calls
            calls += 1
            return canonical_extract(*args, **kwargs)

        monkeypatch.setattr(retrieval_module, "extract_document", counted_extraction)
        assert len(Retriever(config, embedder, manifest, store).search("вибрация")) == 1
        assert calls == 1
        assert cache.load(
            record.document_id,
            record.revision,
            record.extension,
            config.ingestion,
        ) is not None

        # Same size and mtime are insufficient: the source hash is checked before cache use.
        old_stat = source_file.stat()
        source_file.write_bytes(b"X" * old_stat.st_size)
        os.utime(source_file, ns=(old_stat.st_atime_ns, old_stat.st_mtime_ns))
        monkeypatch.setattr(retrieval_module, "extract_document", unexpected_extraction)
        assert Retriever(config, embedder, manifest, store).search("вибрация") == []


def test_reindex_reuses_extracted_cache_when_only_chunk_policy_changes(
    tmp_path, monkeypatch
) -> None:
    config = _test_config(tmp_path)
    source_file = config.paths.source_root / "report.docx"
    source_document = Document()
    source_document.add_paragraph("Длинный расчётный отчёт по вибрации ротора. " * 40)
    source_document.save(source_file)
    embedder = HashingEmbedder(64)

    with (
        ManifestStore(config.manifest_path) as manifest,
        QdrantStore(config.qdrant_path, "index_cache_test", embedder.dimension) as store,
    ):
        assert Indexer(config, embedder, manifest, store).run().indexed == 1
        changed = replace(
            config,
            ingestion=replace(config.ingestion, chunk_chars=400, chunk_overlap_chars=60),
        )

        def unexpected_extraction(*_args, **_kwargs):
            raise AssertionError("reindex must reuse revision-keyed extracted text")

        monkeypatch.setattr(indexer_module, "extract_document", unexpected_extraction)
        report = Indexer(changed, embedder, manifest, store).run()
        assert report.indexed == 1
        assert report.skipped == 0
