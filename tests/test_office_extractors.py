from __future__ import annotations

import logging
import warnings
from dataclasses import replace
from datetime import date
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

import secure_rag.ingestion.extractors as extractors_module
from secure_rag.config import load_config
from secure_rag.ingestion.extractors import ExtractionError, extract_document


@pytest.mark.parametrize(
    ("extension", "helper_name", "logger_namespace", "helper_result"),
    [
        (".pdf", "_extract_pdf", "pypdf", ("Извлечённый текст", 1, False, (1,))),
        (".xlsx", "_extract_xlsx", "openpyxl", ("Извлечённый текст", False)),
    ],
)
def test_parser_warnings_and_logs_do_not_escape_extraction(
    tmp_path,
    monkeypatch,
    capsys,
    caplog,
    extension,
    helper_name,
    logger_namespace,
    helper_result,
) -> None:
    path = tmp_path / f"sample{extension}"
    path.write_bytes(b"placeholder")
    sensitive_diagnostic = "PRIVATE_SOURCE_FRAGMENT_91827"

    def noisy_parser(*_args, **_kwargs):
        warnings.warn(sensitive_diagnostic, UserWarning, stacklevel=1)
        logging.getLogger(f"{logger_namespace}.reader").warning(sensitive_diagnostic)
        return helper_result

    monkeypatch.setattr(extractors_module, helper_name, noisy_parser)
    caplog.clear()
    with caplog.at_level(logging.WARNING):
        result = extract_document(path, load_config().ingestion)

    captured = capsys.readouterr()
    assert result.text == "Извлечённый текст"
    assert sensitive_diagnostic not in captured.out
    assert sensitive_diagnostic not in captured.err
    assert sensitive_diagnostic not in caplog.text

    # Suppression is limited to the parser call and restores the logging setup.
    outside_message = f"{logger_namespace}_OUTSIDE_EXTRACTION"
    with caplog.at_level(logging.WARNING):
        logging.getLogger(f"{logger_namespace}.reader").warning(outside_message)
    assert outside_message in caplog.text


@pytest.mark.parametrize(
    ("extension", "helper_name", "logger_namespace"),
    [
        (".pdf", "_extract_pdf", "pypdf"),
        (".xlsx", "_extract_xlsx", "openpyxl"),
    ],
)
def test_noisy_parser_failures_stay_log_safe(
    tmp_path,
    monkeypatch,
    capsys,
    caplog,
    extension,
    helper_name,
    logger_namespace,
) -> None:
    path = tmp_path / f"broken{extension}"
    path.write_bytes(b"placeholder")
    sensitive_diagnostic = "PRIVATE_FAILURE_FRAGMENT_73465"

    def failing_parser(*_args, **_kwargs):
        warnings.warn(sensitive_diagnostic, UserWarning, stacklevel=1)
        logging.getLogger(f"{logger_namespace}.reader").error(sensitive_diagnostic)
        raise ValueError(sensitive_diagnostic)

    monkeypatch.setattr(extractors_module, helper_name, failing_parser)
    caplog.clear()
    with (
        caplog.at_level(logging.WARNING),
        pytest.raises(ExtractionError) as caught,
    ):
        extract_document(path, load_config().ingestion)

    captured = capsys.readouterr()
    assert caught.value.code == "extract_failed"
    assert str(caught.value) == "extract_failed"
    assert sensitive_diagnostic not in captured.out
    assert sensitive_diagnostic not in captured.err
    assert sensitive_diagnostic not in caplog.text


def test_xlsx_extracts_sheet_names_and_cached_scalar_values(tmp_path) -> None:
    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расчёты"
    sheet.append(["Проект", "Сумма", "Дата", "Активен", "Формула"])
    sheet.append(["Альфа", 42.5, date(2026, 8, 30), True, "=SUM(B2:B2)"])
    second = workbook.create_sheet("Справочник")
    second.append(["Код", 17])
    workbook.save(path)

    config = load_config().ingestion
    first = extract_document(path, config)
    second_result = extract_document(path, config)

    assert first == second_result
    assert "[SHEET] Расчёты" in first.text
    assert "Альфа\t42.5\t2026-08-30\tTRUE" in first.text
    assert "[SHEET] Справочник" in first.text
    assert "Код\t17" in first.text
    assert "=SUM" not in first.text
    assert not first.truncated
    assert [(item.kind, item.value) for item in first.locations] == [
        ("sheet", "Расчёты"),
        ("sheet", "Справочник"),
    ]


def test_pdf_extraction_records_normalized_page_ranges(tmp_path, monkeypatch) -> None:
    path = tmp_path / "pages.pdf"
    path.write_bytes(b"placeholder")
    monkeypatch.setattr(
        extractors_module,
        "_extract_pdf",
        lambda *_args: (
            "первая\n\n[PAGE_BREAK]\n\nвторая",
            2,
            False,
            (1, 2),
        ),
    )

    result = extract_document(path, load_config().ingestion)

    assert result.text == "первая\n\n[PAGE_BREAK]\n\nвторая"
    assert [(item.kind, item.value) for item in result.locations] == [
        ("page", "1"),
        ("page", "2"),
    ]
    assert result.locations[0].end == result.locations[1].start
    assert result.locations[-1].end == len(result.text)


def test_docx_uses_saved_page_count_for_explicitly_approximate_locations(
    tmp_path,
    monkeypatch,
) -> None:
    path = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Первый раздел " * 80)
    document.add_paragraph("Второй раздел " * 80)
    document.save(path)

    saved_result = extract_document(path, load_config().ingestion)
    assert saved_result.page_count == 1
    assert [(item.kind, item.value) for item in saved_result.locations] == [
        ("approx_page", "1")
    ]

    monkeypatch.setattr(extractors_module, "_docx_saved_page_count", lambda _path: 4)
    estimated_result = extract_document(path, load_config().ingestion)

    assert estimated_result.page_count == 4
    assert [item.value for item in estimated_result.locations] == ["1", "2", "3", "4"]
    assert all(item.kind == "approx_page" for item in estimated_result.locations)
    assert estimated_result.locations[0].start == 0
    assert estimated_result.locations[-1].end == len(estimated_result.text)


def test_xlsx_stops_at_extracted_character_limit(tmp_path) -> None:
    path = tmp_path / "bounded.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for index in range(100):
        sheet.append([index, "длинное значение " * 4])
    workbook.save(path)

    config = replace(load_config().ingestion, max_extracted_chars=120)
    result = extract_document(path, config)

    assert len(result.text) <= config.max_extracted_chars
    assert result.truncated


def test_pptx_extracts_text_tables_and_notes_in_slide_order(tmp_path) -> None:
    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    text_box.text = "Заголовок проекта"
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(1))
    table_shape.table.cell(0, 0).text = "Параметр"
    table_shape.table.cell(0, 1).text = "Значение"
    table_shape.table.cell(1, 0).text = "Масса"
    table_shape.table.cell(1, 1).text = "125"
    slide.notes_slide.notes_text_frame.text = "Комментарий докладчика"
    presentation.save(path)

    result = extract_document(path, load_config().ingestion)

    assert result.page_count == 1
    assert result.text.index("[SLIDE 1]") < result.text.index("Заголовок проекта")
    assert "Параметр\tЗначение" in result.text
    assert "Масса\t125" in result.text
    assert "[NOTES]\nКомментарий докладчика" in result.text
    assert not result.truncated
    assert [(item.kind, item.value) for item in result.locations] == [("slide", "1")]


def test_office_archive_with_extreme_compression_is_rejected(tmp_path) -> None:
    path = tmp_path / "archive.xlsx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("xl/worksheets/sheet1.xml", "A" * 1_000_000)

    with pytest.raises(ExtractionError) as caught:
        extract_document(path, load_config().ingestion)

    assert caught.value.code == "office_archive_suspicious_ratio"


def test_docx_uses_the_same_archive_preflight(tmp_path) -> None:
    path = tmp_path / "archive.docx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", "A" * 1_000_000)

    with pytest.raises(ExtractionError) as caught:
        extract_document(path, load_config().ingestion)

    assert caught.value.code == "office_archive_suspicious_ratio"
